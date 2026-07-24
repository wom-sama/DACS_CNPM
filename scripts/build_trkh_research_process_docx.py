from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
USABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEXT = "202124"
MUTED = "5F6368"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D0D7DE"
WHITE = "FFFFFF"
WARNING = "FFF4CE"


def _set_run_font(
    run: Any,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, margins: dict[str, int]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margins[edge]))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell: Any, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_borders(table: Any, color: str = MID_GRAY, size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def _set_table_geometry(table: Any, widths_dxa: Sequence[int]) -> None:
    if sum(widths_dxa) != USABLE_WIDTH_DXA:
        raise ValueError(
            f"Table widths must total {USABLE_WIDTH_DXA} DXA, got {sum(widths_dxa)}"
        )
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths_dxa[index])
            _set_cell_margins(cell, CELL_MARGINS_DXA)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    _set_table_borders(table)


def _set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _prevent_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def _set_paragraph_border(
    paragraph: Any,
    *,
    edge: str,
    color: str,
    size: str = "8",
    space: str = "4",
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = p_bdr.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        p_bdr.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)


def _shade_paragraph(paragraph: Any, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _keep_with_next(paragraph: Any) -> None:
    paragraph.paragraph_format.keep_with_next = True


def _add_hyperlink(paragraph: Any, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Calibri")
    # Keep this minimal. Adding explicit w:sz/w:szCs here made Word and
    # LibreOffice repaginate indefinitely on the living report.
    run_pr.extend([fonts, color, underline])
    run.append(run_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True


def _add_numbering_definition(document: Document, *, bullet: bool) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "\u2022" if bullet else "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Calibri")
    r_pr.append(fonts)
    level.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _apply_numbering(paragraph: Any, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    paragraph.paragraph_format.widow_control = True


def _configure_page(document: Document) -> None:
    # A single page style keeps Word and LibreOffice pagination consistent.
    # LibreOffice 26 can remap explicit even-page parts to a mirrored left-page
    # style, which drops the header and pushes the footer beyond the page edge.
    document.settings.odd_and_even_pages_header_footer = False

    section = document.sections[0]
    title_page = section._sectPr.find(qn("w:titlePg"))
    if title_page is not None:
        section._sectPr.remove(title_page)

    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)


def _add_title_block(document: Document, metadata: dict[str, Any]) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run(metadata["title"].upper())
    _set_run_font(run, size=23, color=TEXT, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run(metadata["subtitle"])
    _set_run_font(run, size=14, color=MUTED)

    rows = [
        ("Nhánh", metadata["branch"]),
        ("Cập nhật", f"{metadata['updated']} | Revision {metadata['revision']}"),
        ("Trạng thái", metadata["status"]),
        ("Phạm vi", metadata["owner"]),
    ]
    for label, value in rows:
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        label_run = p.add_run(f"{label}: ")
        _set_run_font(label_run, size=10.5, color=TEXT, bold=True)
        value_run = p.add_run(value)
        _set_run_font(value_run, size=10.5, color=TEXT)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(10)
    _set_paragraph_border(rule, edge="bottom", color=BLUE, size="12", space="1")


def _add_metric_strip(document: Document, metrics: Sequence[dict[str, str]]) -> None:
    if len(metrics) != 4:
        raise ValueError("Executive metric strip requires exactly four metrics")
    table = document.add_table(rows=1, cols=4)
    _set_repeat_table_header(table.rows[0])
    _set_table_geometry(table, [2340, 2340, 2340, 2340])
    for index, metric in enumerate(metrics):
        cell = table.cell(0, index)
        _set_cell_shading(cell, LIGHT_BLUE if index % 2 == 0 else LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        label = p.add_run(metric["label"].upper())
        _set_run_font(label, size=8.2, color=MUTED, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(1)
        value = p2.add_run(metric["value"])
        _set_run_font(value, size=14, color=DARK_BLUE, bold=True)
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(0)
        p3.paragraph_format.line_spacing = 1.0
        detail = p3.add_run(metric["detail"])
        _set_run_font(detail, size=8.5, color=TEXT)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_callout(document: Document, label: str, text: str, *, warning: bool = False) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Pt(9)
    p.paragraph_format.right_indent = Pt(8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    _shade_paragraph(p, WARNING if warning else LIGHT_BLUE)
    _set_paragraph_border(p, edge="left", color=BLUE, size="18", space="6")
    label_run = p.add_run(f"{label}: ")
    _set_run_font(label_run, size=10.5, color=DARK_BLUE, bold=True)
    text_run = p.add_run(text)
    _set_run_font(text_run, size=10.5, color=TEXT)


def _add_caption(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    _set_run_font(run, size=9, color=MUTED, italic=True)


def _add_table(document: Document, block: dict[str, Any]) -> None:
    headers = block["headers"]
    rows = block["rows"]
    widths = block["widths_dxa"]
    if len(headers) != len(widths):
        raise ValueError(f"Header/width mismatch for {block.get('caption', 'table')}")
    for row in rows:
        if len(row) != len(headers):
            raise ValueError(f"Row/header mismatch for {block.get('caption', 'table')}")
    _add_caption(document, block["caption"])
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_row = table.rows[0]
    for index, value in enumerate(headers):
        cell = header_row.cells[index]
        _set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
            if len(value) < 18
            else WD_ALIGN_PARAGRAPH.LEFT
        )
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(value)
        _set_run_font(run, size=9, color=TEXT, bold=True)
    _set_repeat_table_header(header_row)

    for row_values in rows:
        row = table.add_row()
        _prevent_row_split(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            p = cell.paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index > 0 and len(value) <= 24
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            p.paragraph_format.widow_control = True
            run = p.add_run(value)
            _set_run_font(run, size=9, color=TEXT)
    _set_table_geometry(table, widths)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def _add_list(
    document: Document,
    items: Iterable[str],
    *,
    num_id: int,
) -> None:
    for item in items:
        p = document.add_paragraph()
        _apply_numbering(p, num_id)
        run = p.add_run(item)
        _set_run_font(run, size=11, color=TEXT)


def _add_sources(document: Document, items: Sequence[dict[str, str]], num_id: int) -> None:
    for item in items:
        p = document.add_paragraph()
        p.paragraph_format.keep_together = True
        _apply_numbering(p, num_id)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(1)
        title = p.add_run(item["title"])
        _set_run_font(title, size=10.5, color=TEXT, bold=True)
        detail = p.add_run(f" - {item['detail']} ")
        _set_run_font(detail, size=10.5, color=TEXT)
        _add_hyperlink(p, "Mở nguồn", item["url"])


def _add_block(
    document: Document,
    block: dict[str, Any],
    *,
    bullet_num_id: int,
    decimal_num_id: int,
) -> None:
    kind = block["type"]
    if kind == "paragraph":
        p = document.add_paragraph()
        p.paragraph_format.widow_control = True
        run = p.add_run(block["text"])
        _set_run_font(run, size=11, color=TEXT)
    elif kind == "callout":
        _add_callout(document, block["label"], block["text"])
    elif kind == "note":
        _add_callout(document, "Lưu ý", block["text"], warning=True)
    elif kind == "bullets":
        _add_list(document, block["items"], num_id=bullet_num_id)
    elif kind == "numbered":
        block_num_id = _add_numbering_definition(document, bullet=False)
        _add_list(document, block["items"], num_id=block_num_id)
    elif kind == "table":
        _add_table(document, block)
    elif kind == "sources":
        source_num_id = _add_numbering_definition(document, bullet=False)
        _add_sources(document, block["items"], source_num_id)
    elif kind == "page_break":
        p = document.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
    else:
        raise ValueError(f"Unsupported block type: {kind}")


def _audit_source(payload: dict[str, Any]) -> None:
    document = payload.get("document")
    sections = payload.get("sections")
    if not isinstance(document, dict) or not isinstance(sections, list) or not sections:
        raise ValueError("Source must contain document metadata and non-empty sections")
    required_meta = {
        "title",
        "subtitle",
        "revision",
        "updated",
        "branch",
        "status",
        "owner",
    }
    missing = required_meta - set(document)
    if missing:
        raise ValueError(f"Missing document metadata: {sorted(missing)}")
    if len(payload.get("executive_metrics", [])) != 4:
        raise ValueError("Exactly four executive metrics are required")
    for section in sections:
        if not section.get("title") or not section.get("blocks"):
            raise ValueError("Every section requires title and blocks")
        for block in section["blocks"]:
            if block["type"] == "table":
                if sum(block["widths_dxa"]) != USABLE_WIDTH_DXA:
                    raise ValueError(
                        f"Table {block.get('caption')} widths do not total "
                        f"{USABLE_WIDTH_DXA} DXA"
                    )
            if block["type"] in {"bullets", "numbered"}:
                for item in block["items"]:
                    if item.startswith(("-", "*", "\u2022")):
                        raise ValueError("List items must not contain fake markers")


def build_document(source: Path, output: Path) -> None:
    payload = json.loads(source.read_text(encoding="utf-8"))
    _audit_source(payload)

    document = Document()
    _configure_page(document)
    _configure_styles(document)
    bullet_num_id = _add_numbering_definition(document, bullet=True)
    decimal_num_id = _add_numbering_definition(document, bullet=False)
    metadata = payload["document"]
    _add_title_block(document, metadata)
    _add_metric_strip(document, payload["executive_metrics"])

    for section in payload["sections"]:
        heading = document.add_paragraph(style="Heading 1")
        heading.add_run(section["title"])
        _keep_with_next(heading)
        for block in section["blocks"]:
            _add_block(
                document,
                block,
                bullet_num_id=bullet_num_id,
                decimal_num_id=decimal_num_id,
            )

    properties = document.core_properties
    properties.title = f"{metadata['title']} - {metadata['subtitle']}"
    properties.subject = "TRKH 5-class research process and status"
    properties.author = "TRKH classification-only research"
    properties.keywords = "TRKH, classification, CNN, Transformer, XAI, synthetic data"
    properties.comments = (
        f"Generated from {source.name}; revision {metadata['revision']}; "
        "standard_business_brief + memo_masthead"
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Wrote {output}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the reproducible TRKH research process/status DOCX."
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("docs/TRKH_5CLASS_RESEARCH_PROCESS_STATUS.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/TRKH_5CLASS_RESEARCH_PROCESS_STATUS.docx"),
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build_document(args.source.resolve(), args.output.resolve())
